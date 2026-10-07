"""Hito 2AZ, fase 5: muestra v2 para contraste manual (10 facturas PIO).

Uso: python -B pruebas/auditoria_2az/muestra_v2.py <motor_banco.json> <salida.xlsx>

Ejecuta el motor local vigente (en memoria) sobre los PDF seleccionados del mirror
(solo lectura; ruta localizada por SHA en el banco) y vuelca, por factura:
vencimientos (fecha e importe extraidos, R14), recargo de equivalencia total y por
tipo, y movimientos (cargos de servicio y abonos) con importe, mas columnas vacias
para el valor real. R13: ninguna columna se deriva del nombre del archivo; la hoja
"Localizacion" relaciona SHA y ruta solo para abrir el PDF.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "pruebas/auditoria_2ay"))

from banco_extraccion import _salida_fuera_del_repo  # noqa: E402
from src.facturas.motor_local.backend.pdfium import BackendPdfium  # noqa: E402
from src.facturas.motor_local.servicio import MotorDocumentoLocal  # noqa: E402

ALLIANCE = ("08007971", "08007973", "08007969", "08011733", "08011736", "08007501", "08009716")
OTROS_LAYOUTS = ("hefame-local", "cofares-local", "fedefarma-local")


def _v(campo):
    return campo.get("valor") if isinstance(campo, dict) else None


def _seleccion(docs: list[dict]) -> list[tuple[dict, str]]:
    elegidas = []
    for numero in ALLIANCE:
        d = next(d for d in docs for f in (d.get("oficial") or {}).get("facturas_local") or [] if f.get("numero") == numero)
        elegidas.append((d, numero))
    for layout in OTROS_LAYOUTS:
        d = next(d for d in docs if (d.get("oficial") or {}).get("layout") == layout
                 and (d.get("oficial") or {}).get("documento_completo_demostrado") is True)
        elegidas.append((d, d["oficial"]["facturas_local"][0]["numero"]))
    return elegidas


def _factura_local(local, numero):
    for f in local.facturas:
        cab = f.get("cabecera") or local.cabecera or {}
        if _v(cab.get("numero_factura")) == numero:
            return f, cab
    return {}, local.cabecera or {}


def _fila(d: dict, numero: str) -> dict:
    local = MotorDocumentoLocal(BackendPdfium()).extraer(d["ruta"])
    f, cab = _factura_local(local, numero)
    impuestos = f.get("impuestos") or local.impuestos or []
    vencimientos = f.get("vencimientos") or local.vencimientos or []
    # Una sola entrada por fecha (las repeticiones de hoja son evidencia del mismo vencimiento).
    por_fecha: dict = {}
    for v in vencimientos:
        fecha = _v(v.get("fecha"))
        if fecha is not None and (fecha not in por_fecha or por_fecha[fecha] is None):
            por_fecha[fecha] = _v(v.get("importe"))
    re_por_tipo = [f"{_v(i.get('tipo_recargo_equivalencia'))}% -> {_v(i.get('cuota_recargo_equivalencia'))}"
                   for i in impuestos if _v(i.get("cuota_recargo_equivalencia")) not in (None, 0, 0.0)]
    movimientos = []
    if local.documento.get("layout") == "alliance-local":
        # Lo que se persiste: movimientos_comerciales del puente (operaciones + movimientos no relacionados).
        from src.facturas.runtime_supabase.multifactura import adaptar_resultado_local
        normal = next(x for x in adaptar_resultado_local(local)["facturas"] if _v(x["numero_factura"]) == numero)
        for m in normal["movimientos_comerciales"]:
            movimientos.append(f"{_v(m.get('descripcion_literal'))} [{m.get('tipo')}/{m.get('sentido')}] "
                               f"{_v(m.get('importe'))}")
    else:
        for m in f.get("movimientos") or local.movimientos or []:
            movimientos.append(f"{_v(m.get('descripcion_literal'))} [{m.get('categoria')}/{m.get('sentido')}] "
                               f"{_v(m.get('importe'))}")
    return {
        "documento_sha256": d["archivo_hash"], "documento_id": d["id"],
        "paginas": (f.get("segmento") or {}).get("paginas"), "layout": local.documento.get("layout"),
        "proveedor": _v((cab.get("proveedor") or {}).get("nombre")), "numero": numero,
        "fecha": _v(cab.get("fecha_factura")), "base": _v(cab.get("base_imponible_total")),
        "iva": _v(cab.get("iva_total")), "re_total": _v(cab.get("recargo_equivalencia_total")),
        "re_por_tipo": "; ".join(re_por_tipo) or None, "total": _v(cab.get("importe_total")),
        "vencimientos": "; ".join(f"{k}: {v if v is not None else 'SIN IMPORTE'}" for k, v in por_fecha.items()),
        "movimientos": " | ".join(movimientos) or None, "ruta": d["ruta"],
    }


def main(motor: Path, salida: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    docs = json.loads(motor.read_text(encoding="utf-8"))["documentos"]
    filas = [_fila(d, numero) for d, numero in _seleccion(docs)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Muestra v2"
    columnas = [("documento_sha256", "SHA-256 del PDF"), ("paginas", "Páginas (factura)"), ("layout", "Layout"),
                ("proveedor", "Proveedor (contenido)"), ("numero", "Número"), ("fecha", "Fecha"), ("base", "Base"),
                ("iva", "IVA"), ("re_total", "RE total"), ("re_por_tipo", "RE por tipo"), ("total", "Total"),
                ("vencimientos", "Vencimientos (fecha: importe)"), ("movimientos", "Movimientos (cargos servicio / abonos)")]
    vacias = ["Fecha venc. REAL", "Importe venc. REAL", "RE total REAL", "RE por tipo REAL", "Movimientos REAL",
              "¿Coincide?", "Observaciones de Pío"]
    ws.append([t for _, t in columnas] + vacias)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for celda in ws[1][len(columnas):]:
        celda.fill = PatternFill("solid", fgColor="FFF2CC")
    for fila in filas:
        ws.append([str(fila[k]) if isinstance(fila[k], list) else fila[k] for k, _ in columnas] + [None] * len(vacias))
    for i, ancho in enumerate([20, 10, 18, 30, 14, 12, 11, 10, 10, 26, 11, 34, 70] + [16] * len(vacias), start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = ancho
    ws.freeze_panes = "B2"
    loc = wb.create_sheet("Localizacion")
    loc.append(["SHA-256 del PDF", "Ruta en el mirror (solo para abrir el PDF; no es fuente de datos, R13)"])
    for fila in filas:
        loc.append([fila["documento_sha256"], fila["ruta"]])
    wb.save(salida)
    print(json.dumps([{k: v for k, v in f.items() if k != "ruta"} for f in filas], ensure_ascii=False, indent=1,
                     default=str))


if __name__ == "__main__":
    main(Path(sys.argv[1]), _salida_fuera_del_repo(sys.argv[2]))
