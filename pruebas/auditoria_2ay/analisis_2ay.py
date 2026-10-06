"""Hito 2AY: metricas (fase 3), simulacion de conciliacion (fase 4) y muestra de contraste (fase 5).

Uso:
  python -B pruebas/auditoria_2ay/analisis_2ay.py <supabase.json> <inventario.json> <motor.json> <directorio_salida>

No se conecta a nada: trabaja sobre las salidas READ_ONLY de ``supabase_readonly.py``
y ``banco_extraccion.py``. La conciliacion se simula en memoria con las
funciones oficiales (``buscar_candidato_albaran``, ``detalle_desde_busqueda``,
``construir_detalles_movimientos_documentales``, ``conciliar_importes``) contra
los albaranes leidos de Supabase; la elegibilidad replica en Python la funcion
``cf_evaluar_elegibilidad_conciliacion`` (migracion 14). Sin claim ni escritura.
Escribe metricas.json, tablas.md, diferencias_vs_albaranes.svg y
muestra_contraste.xlsx en el directorio indicado (fuera del repositorio).
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from src.facturas.runtime_supabase.conciliacion import (  # noqa: E402
    VENTANA_FECHA_DIAS, AlbaranDocumentalTrabajo, AjusteDocumentalTrabajo, CandidatoAlbaranSupabase,
    MovimientoDocumentalTrabajo, buscar_candidato_albaran, conciliar_importes,
    construir_detalles_movimientos_documentales, detalle_desde_busqueda)

T1 = Decimal("0.05")
NIF_PIO = "40901058C"
ESTADOS_OK = {"VALIDADA", "VALIDADA_CON_INCIDENCIAS"}


def _d(valor) -> Decimal | None:
    if valor in (None, ""):
        return None
    return Decimal(str(valor))


def tolerancias(casados: int) -> dict[str, Decimal]:
    """T1 fija 0,05; T2/T3: 0,01 EUR por albaran casado con tope 1,00 / 0,50."""
    return {"T1": T1, "T2": min(Decimal("0.01") * casados, Decimal("1.00")),
            "T3": min(Decimal("0.01") * casados, Decimal("0.50"))}


# --------------------------------------------------------------------------- elegibilidad (replica migracion 14)

def elegibilidad(f: dict, tolerancia: Decimal = T1) -> tuple[str, str]:
    """Replica de cf_evaluar_elegibilidad_conciliacion sobre una factura normalizada autorizable."""
    nif = re.sub(r"[^A-Z0-9]", "", str(f.get("destinatario_nif") or "").upper())
    nombre = unicodedata.normalize("NFKD", str(f.get("destinatario_nombre") or "")).encode("ascii", "ignore") \
        .decode().upper()
    if not f.get("destinatario_con_evidencia"):
        return "NO_APTA", "FARMACIA_NO_DEMOSTRABLE"
    if not (nif == NIF_PIO or (nif == "" and re.search(r"PUIG SALOM.*PIO", nombre))):
        return "NO_APTA", "FARMACIA_NO_CONSISTENTE"
    if f.get("incidencias_bloqueantes"):
        return "NO_APTA", "INCIDENCIA_BLOQUEANTE"
    tipo = {"MERCANCIA": "M", "SERVICIOS": "S", "MIXTA": "X"}.get(f.get("naturaleza_principal"))
    if tipo is None:
        return "REQUIERE_REVISION", "TIPO_DOCUMENTAL_NO_DEMOSTRADO"
    total = _d(f.get("total"))
    if total is None:
        return "NO_APTA", "TOTAL_NO_DEMOSTRADO"
    albaranes = [a for a in f.get("albaranes", []) if _d(a.get("importe_total")) is not None]
    suma_alb = sum((-abs(_d(a["importe_total"])) if a.get("sentido") == "ABONO" else abs(_d(a["importe_total"]))
                    for a in albaranes), Decimal(0))
    movs = f.get("movimientos", [])
    incompletos = sum(1 for m in movs if not str(m.get("descripcion") or "").strip() or m.get("sentido") is None
                      or (m.get("importe") is None and m.get("base") is None))
    con_importe = sum(1 for m in movs if m.get("importe") is not None)
    con_base = sum(1 for m in movs if m.get("base") is not None)
    signo = lambda m, k: -abs(_d(m[k])) if m.get("sentido") == "ABONO" else abs(_d(m[k]))  # noqa: E731
    suma_imp = sum((signo(m, "importe") for m in movs if m.get("importe") is not None), Decimal(0))
    suma_base = sum((signo(m, "base") for m in movs if m.get("base") is not None), Decimal(0))
    base = _d(f.get("base_imponible"))
    fiscal = base is not None and abs(total - (base + (_d(f.get("iva")) or 0) + (_d(f.get("recargo_equivalencia")) or 0)
                                                + (_d(f.get("otros")) or 0))) <= tolerancia
    if tipo == "M":
        return ("APTA", "APTA_MERCANCIA") if albaranes else ("NO_APTA", "FALTAN_ALBARANES_MERCANCIA")
    traza = len(movs) > 0 and incompletos == 0
    if tipo == "S":
        explicado = ((con_importe == len(movs) and abs(total - suma_imp) <= tolerancia)
                     or (con_importe == 0 and con_base == len(movs) and base is not None
                         and abs(base - suma_base) <= tolerancia and fiscal))
        if not traza:
            return "NO_APTA", "TRAZABILIDAD_SERVICIO_INSUFICIENTE"
        if not fiscal:
            return "NO_APTA", "FISCALIDAD_INCOHERENTE"
        return ("APTA", "APTA_GASTO_SERVICIO") if explicado else ("NO_APTA", "TOTAL_NO_EXPLICADO")
    if not albaranes:
        return "NO_APTA", "FALTAN_ALBARANES_MERCANCIA"
    if not traza or con_importe != len(movs):
        return "NO_APTA", "TRAZABILIDAD_SERVICIO_INSUFICIENTE"
    if not fiscal:
        return "NO_APTA", "FISCALIDAD_INCOHERENTE"
    if abs(total - (suma_alb + suma_imp)) > tolerancia:
        return "NO_APTA", "TOTAL_NO_EXPLICADO"
    return "APTA", "APTA_MIXTA"


# --------------------------------------------------------------------------- conciliacion simulada

def candidatos_operacionales(albaranes: list[dict]) -> list[CandidatoAlbaranSupabase]:
    return [CandidatoAlbaranSupabase(
        id_contador=int(a["id_contador"]), farmacia=a["farmacia"], id_proveedor=str(a["id_proveedor"]),
        proveedor=str(a["proveedor"]), numero_albaran=str(a["numero_albaran"]), fecha=date.fromisoformat(a["fecha"]),
        importe_puc=_d(a.get("importe_puc")), importe_pvp=_d(a.get("importe_pvp")), estado=a.get("estado"))
        for a in albaranes]


def simular_conciliacion(f: dict, operacionales: list[CandidatoAlbaranSupabase]) -> dict:
    """Equivalente en memoria de construir_detalles + conciliar_importes (sin claim)."""
    naturaleza = f.get("naturaleza_principal")
    usa_mercancia, usa_servicios = naturaleza in {"MERCANCIA", "MIXTA"}, naturaleza in {"SERVICIOS", "MIXTA"}
    if not usa_mercancia and not usa_servicios:
        return {"camino": "FALLO", "codigo": "TIPO_DOCUMENTAL_NO_DEMOSTRADO"}
    detalles = []
    if usa_servicios:
        movs = [MovimientoDocumentalTrabajo(id=f"mov-{i}", concepto_literal=str(m.get("descripcion")),
                                            tipo=str(m.get("tipo")), sentido=m.get("sentido"),
                                            importe=_d(m.get("importe")), base=_d(m.get("base")))
                for i, m in enumerate(f.get("movimientos", []), 1)]
        ajustes = ()
        if naturaleza != "MIXTA" and movs and all(m.importe is None for m in movs):
            ajustes = tuple(AjusteDocumentalTrabajo(n, _d(v), {}) for n, v in
                            (("IVA_TOTAL", f.get("iva")), ("RECARGO_EQUIVALENCIA_TOTAL", f.get("recargo_equivalencia")))
                            if _d(v) not in (None, Decimal(0)))
        movimiento_detalles = construir_detalles_movimientos_documentales(movs, ajustes=ajustes)
        if movimiento_detalles is None:
            return {"camino": "FALLO", "codigo": "TRAZABILIDAD_SERVICIO_INSUFICIENTE"}
        detalles.extend(movimiento_detalles)
    documentales = []
    if usa_mercancia:
        documentales = [AlbaranDocumentalTrabajo(
            id=f"alb-{i}", numero=a.get("numero"), fecha=date.fromisoformat(a["fecha"]) if a.get("fecha") else None,
            importe=_d(a.get("importe_total")), sentido=a.get("sentido")) for i, a in enumerate(f.get("albaranes", []), 1)]
    fechas = [a.fecha for a in documentales if a.fecha]
    ventana = []
    if fechas:
        inicio, fin = min(fechas) - timedelta(days=VENTANA_FECHA_DIAS), max(fechas) + timedelta(days=VENTANA_FECHA_DIAS)
        ventana = [c for c in operacionales if inicio <= c.fecha <= fin]
    matching = Counter()
    for doc in documentales:
        busqueda = buscar_candidato_albaran(doc, ventana, proveedor_literal=f.get("proveedor_nombre"))
        matching[busqueda.estado] += 1
        detalles.append(detalle_desde_busqueda(doc, busqueda))
    resultado = conciliar_importes(_d(f.get("total")), detalles, tolerancia=T1)
    casados = sum(1 for d in detalles if d.tipo_relacion.value == "UNO_A_UNO")
    return {"camino": "PERSISTIR", "resultado_t1": resultado.resultado, "importe_explicado": str(resultado.importe_explicado),
            "diferencia": str(resultado.diferencia), "albaranes_documentales": len(documentales),
            "albaranes_casados": casados, "albaranes_sin_coincidencia": len(documentales) - casados,
            "movimientos": len(f.get("movimientos", [])), "matching": dict(matching)}


# --------------------------------------------------------------------------- utilidades de informe

def _proveedor(fila: dict) -> str:
    return (fila.get("oficial") or {}).get("layout") or "SIN_LAYOUT"


def _tabla(cabecera: list[str], filas: list[list]) -> str:
    salida = ["| " + " | ".join(cabecera) + " |", "|" + "|".join("---" for _ in cabecera) + "|"]
    salida += ["| " + " | ".join("" if v is None else str(v) for v in f) + " |" for f in filas]
    return "\n".join(salida)


def _svg_dispersion(puntos: list[tuple[int, float, str]], destino: Path) -> None:
    """Grafico SVG sin dependencias: diferencia (EUR, |x|) frente a albaranes casados, con T1/T2/T3."""
    ancho, alto, m = 760, 420, 60
    max_x = max([p[0] for p in puntos] + [10])
    max_y = max([abs(p[1]) for p in puntos] + [1.0])
    max_y = min(max_y, 5.0)
    sx = lambda x: m + x / max_x * (ancho - 2 * m)  # noqa: E731
    sy = lambda y: alto - m - min(abs(y), max_y) / max_y * (alto - 2 * m)  # noqa: E731
    e = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{ancho}" height="{alto}" font-family="sans-serif" font-size="11">',
         f'<rect width="{ancho}" height="{alto}" fill="white"/>',
         f'<line x1="{m}" y1="{alto-m}" x2="{ancho-m}" y2="{alto-m}" stroke="#444"/>',
         f'<line x1="{m}" y1="{m}" x2="{m}" y2="{alto-m}" stroke="#444"/>',
         f'<text x="{ancho/2}" y="{alto-20}" text-anchor="middle">Albaranes casados 1:1</text>',
         f'<text x="16" y="{alto/2}" transform="rotate(-90 16 {alto/2})" text-anchor="middle">|diferencia| EUR (tope {max_y})</text>']
    for i in range(0, max_x + 1, max(1, max_x // 10)):
        e.append(f'<text x="{sx(i)}" y="{alto-m+14}" text-anchor="middle">{i}</text>')
    for j in range(6):
        y = max_y * j / 5
        e.append(f'<text x="{m-6}" y="{sy(y)+4}" text-anchor="end">{y:.2f}</text>')
    e.append(f'<line x1="{sx(0)}" y1="{sy(0.05)}" x2="{sx(max_x)}" y2="{sy(0.05)}" stroke="#1f77b4" stroke-dasharray="4"/>')
    for tope, color in ((1.0, "#d62728"), (0.5, "#2ca02c")):
        x_tope = min(max_x, tope / 0.01)
        e.append(f'<polyline fill="none" stroke="{color}" stroke-dasharray="2" points="{sx(0)},{sy(0)} '
                 f'{sx(x_tope)},{sy(tope)} {sx(max_x)},{sy(tope)}"/>')
    for x, y, etiqueta in puntos:
        e.append(f'<circle cx="{sx(x)}" cy="{sy(y)}" r="4" fill="#ff7f0e" fill-opacity="0.8"><title>{etiqueta}: '
                 f'{x} casados, dif {y:.2f}</title></circle>')
    e.append(f'<text x="{ancho-m}" y="{m-20}" text-anchor="end">azul T1 0,05 · rojo T2 0,01/alb tope 1,00 · '
             f'verde T3 tope 0,50</text>')
    e.append("</svg>")
    destino.write_text("\n".join(e), encoding="utf-8")


# --------------------------------------------------------------------------- principal

def main(supabase_ruta: Path, inventario_ruta: Path, motor_ruta: Path, salida: Path) -> None:
    supabase = json.loads(supabase_ruta.read_text(encoding="utf-8"))
    inv = {d["id"]: d for d in json.loads(inventario_ruta.read_text(encoding="utf-8"))["documentos"]}
    docs = json.loads(motor_ruta.read_text(encoding="utf-8"))["documentos"]
    operacionales = candidatos_operacionales(supabase["albaranes"])
    produccion_por_documento = defaultdict(list)
    for f in supabase["facturas"]:
        produccion_por_documento[f["documento_id"]].append(f)
    m: dict = {}

    # 3.1 embudo global
    con_texto = [d for d in docs if inv[d["id"]].get("capa_texto") in {"TEXTO", "MIXTO"}]
    reconocidos = [d for d in docs if (d.get("oficial") or {}).get("layout")]
    completos = [d for d in reconocidos if d["oficial"].get("documento_completo_demostrado") is True]
    con_autoridad = [d for d in completos if d["oficial"].get("etapa_detencion") not in {"AUTORIDAD"}]
    autorizados = [d for d in docs if d["oficial"].get("segmentos_autorizados")]
    facturas_autorizables = sum(len(d["oficial"]["segmentos_autorizados"]) for d in autorizados)
    m["embudo"] = {
        "documentos": len(docs),
        "localizados_en_mirror": sum(1 for d in docs if d.get("ruta")),
        "con_texto": len(con_texto),
        "proveedor_reconocido": len(reconocidos),
        "extraccion_completa_demostrada": len(completos),
        "con_autoridad_manual": len(con_autoridad),
        "documentos_con_facturas_autorizables": len(autorizados),
        "facturas_detectadas_en_reconocidos": sum(d["oficial"].get("facturas_detectadas", 0) for d in reconocidos),
        "facturas_autorizables": facturas_autorizables,
    }
    m["etapas_detencion"] = Counter(f'{d["oficial"].get("etapa_detencion")}:{d["oficial"].get("motivo")}' for d in docs)

    # 3.2 por proveedor
    por_proveedor: dict = {}
    for prov in sorted({_proveedor(d) for d in docs}):
        grupo = [d for d in docs if _proveedor(d) == prov]
        facturas = sum(d["oficial"].get("facturas_detectadas", 0) for d in grupo)
        autorizables = sum(len(d["oficial"].get("segmentos_autorizados") or []) for d in grupo)
        motivos = Counter()
        for d in grupo:
            o = d["oficial"]
            if o.get("etapa_detencion") != "PREVIO_A_PERSISTIR_OK":
                motivos[f'{o.get("etapa_detencion")}:{o.get("motivo")}'] += 1
            for f in o.get("facturas_normalizadas") or []:
                for motivo in f.get("motivos_no_autorizable") or []:
                    motivos[f"FACTURA:{motivo}"] += 1
        cuadran = sum(1 for d in grupo for f in d["oficial"].get("facturas_local") or []
                      if (f.get("cuadre_fiscal") or {}).get("cuadra_005"))
        completos_g = sum(1 for d in grupo if d["oficial"].get("documento_completo_demostrado") is True)
        por_proveedor[prov] = {"documentos": len(grupo), "completos": completos_g, "facturas": facturas,
                               "facturas_con_cuadre_fiscal": cuadran, "autorizables": autorizables,
                               "pct_autorizables": round(100 * autorizables / facturas, 1) if facturas else 0.0,
                               "motivos_principales": motivos.most_common(3)}
    m["por_proveedor"] = por_proveedor

    # 3.3 sin texto
    sin_texto = []
    for d in docs:
        i = inv[d["id"]]
        if i.get("capa_texto") in {"SIN_TEXTO", "MIXTO"}:
            nombre = Path(d.get("ruta") or i.get("archivo_nombre") or "").name
            sin_texto.append({"archivo": nombre, "capa_texto": i.get("capa_texto"), "paginas": i.get("paginas"),
                              "proveedor_probable_nombre": re.split(r"[\s_.]+", nombre.upper())[0] if nombre else None,
                              "metadatos": {k: v for k, v in (i.get("metadatos_pdf") or {}).items()
                                            if k in ("Producer", "Creator", "Author", "Title")},
                              "etapa": d["oficial"].get("etapa_detencion"), "motivo": d["oficial"].get("motivo")})
    m["sin_texto"] = sin_texto

    # 3.4 no reconocidos con texto (y excepciones del motor)
    grupos = defaultdict(list)
    for d in docs:
        o = d["oficial"]
        if o.get("layout") or inv[d["id"]].get("capa_texto") not in {"TEXTO", "MIXTO"}:
            continue
        i = inv[d["id"]]
        cifs = [c for c in i.get("cifs") or []]
        clave = f"CIF:{cifs[0]}" if cifs else "CABECERA:" + " / ".join((i.get("cabecera_primera_pagina") or [])[:2])[:60]
        grupos[clave].append({"archivo": Path(d.get("ruta") or "").name, "etapa": o.get("etapa_detencion"),
                              "motivo": o.get("motivo"), "excepcion": (o.get("excepcion") or {}).get("tipo"),
                              "todos_cifs": cifs[:4]})
    m["no_reconocidos_con_texto"] = {
        "documentos": sum(len(v) for v in grupos.values()), "grupos": len(grupos),
        "grupos_detalle": sorted(([k, len(v), v] for k, v in grupos.items()), key=lambda x: -x[1])}
    excepciones = Counter()
    for d in docs:
        e = d["oficial"].get("excepcion")
        if e and d["oficial"].get("etapa_detencion") == "EXTRACCION_MOTOR":
            excepciones[f'{e.get("tipo")} @ {e.get("traza")}'] += 1
    m["excepciones_motor"] = excepciones.most_common()

    # 3.5 Alliance simulacion 2AU
    alliance = []
    for d in docs:
        if _proveedor(d) != "alliance-local":
            continue
        oficial = {f["numero"]: f for f in d["oficial"].get("facturas_normalizadas") or []}
        simulado = {f["numero"]: f for f in (d.get("simulado_2au") or {}).get("facturas_normalizadas") or []}
        locales = {f["numero"]: f for f in d["oficial"].get("facturas_local") or []}
        for numero, f in oficial.items():
            s = simulado.get(numero) or {}
            loc = locales.get(numero) or {}
            alliance.append({
                "documento": Path(d["ruta"]).name, "numero": numero, "total": f.get("total"),
                "autorizable_oficial": f.get("autorizable"), "autorizable_2au": s.get("autorizable"),
                "albaranes_oficial": len(f.get("albaranes", [])), "albaranes_2au": len(s.get("albaranes", [])),
                "movimientos_oficial": len(f.get("movimientos", [])), "movimientos_2au": len(s.get("movimientos", [])),
                "naturaleza_oficial": f.get("naturaleza_principal"), "naturaleza_2au": s.get("naturaleza_principal"),
                "filas_no_promovidas": [(x["tipo_pedido"], x["sentido"], x["total"]) for x in loc.get("filas_no_promovidas", [])],
                "elegibilidad_oficial": elegibilidad(f) if f.get("autorizable") else None,
                "elegibilidad_2au": elegibilidad(s) if s.get("autorizable") else None,
            })
    m["alliance_2au"] = {
        "facturas": len(alliance),
        "con_filas_no_promovidas": sum(1 for a in alliance if a["filas_no_promovidas"]),
        "tipos_no_promovidos": Counter(t for a in alliance for t, _, _ in a["filas_no_promovidas"]),
        "autorizables_oficial": sum(1 for a in alliance if a["autorizable_oficial"]),
        "autorizables_2au": sum(1 for a in alliance if a["autorizable_2au"]),
        "aptas_conciliacion_oficial": sum(1 for a in alliance if (a["elegibilidad_oficial"] or ("",))[0] == "APTA"),
        "aptas_conciliacion_2au": sum(1 for a in alliance if (a["elegibilidad_2au"] or ("",))[0] == "APTA"),
        "cambian": [a for a in alliance if (a["albaranes_oficial"], a["movimientos_oficial"], a["elegibilidad_oficial"])
                    != (a["albaranes_2au"], a["movimientos_2au"], a["elegibilidad_2au"])],
        "detalle": alliance,
    }

    # 3.6 coherencia con produccion
    coherencia = []
    por_id = {d["id"]: d for d in docs}
    for documento_id, facturas in produccion_por_documento.items():
        d = por_id.get(documento_id)
        for p in facturas:
            fila = {"numero": p["numero_factura"], "documento": documento_id[:8], "proveedor": p["proveedor_literal"]}
            if d is None or not d.get("ruta"):
                fila["resultado"] = "NO_EVALUABLE_DOCUMENTO_AUSENTE_DEL_MIRROR"
                coherencia.append(fila)
                continue
            o = d["oficial"]
            normal = next((f for f in o.get("facturas_normalizadas") or [] if f["numero"] == p["numero_factura"]), None)
            local = next((f for f in o.get("facturas_local") or [] if f.get("numero") == p["numero_factura"]), None)
            fuente = normal or local
            fila["layout"] = o.get("layout")
            fila["etapa_motor"] = o.get("etapa_detencion")
            if fuente is None:
                fila["resultado"] = "NO_REPRODUCIDA"
                fila["numeros_locales"] = [f.get("numero") for f in o.get("facturas_local") or []]
                coherencia.append(fila)
                continue
            total_local = fuente.get("total")
            comparacion = {
                "numero": [p["numero_factura"], fuente.get("numero")],
                "fecha": [p["fecha_factura"], fuente.get("fecha")],
                "total": [p["importe_total"], total_local],
                "base": [p["base_imponible_total"], fuente.get("base_imponible", fuente.get("base"))],
                "iva": [p["iva_total"], fuente.get("iva")],
            }
            diferencias = {k: v for k, v in comparacion.items()
                           if (_d(v[0]) != _d(v[1]) if k in {"total", "base", "iva"} else str(v[0]) != str(v[1]))}
            if normal is not None:
                identidad = [p["identidad_economica_clave"], normal.get("identidad_economica_clave")]
                if identidad[0] != identidad[1]:
                    diferencias["identidad_economica_clave"] = identidad
                fila["identidad"] = "IGUAL" if identidad[0] == identidad[1] else "DISTINTA"
            else:
                fila["identidad"] = "NO_COMPARABLE_SIN_PUENTE_CERTIFICADO"
            fila["fuente_local"] = "NORMALIZADA" if normal else "EXTRACCION_LOCAL"
            fila["diferencias"] = diferencias
            fila["resultado"] = "REPRODUCE" if not diferencias else "DIFIERE"
            coherencia.append(fila)
    m["coherencia_produccion"] = coherencia

    # 4 simulacion de conciliacion
    produccion_numeros = {f["numero_factura"]: f for f in supabase["facturas"]}
    simulaciones, puntos = [], []
    for d in docs:
        for origen in ("oficial", "simulado_2au"):
            for f in (d.get(origen) or {}).get("facturas_normalizadas") or []:
                if not f.get("autorizable"):
                    continue
                if origen == "simulado_2au":
                    oficial_f = next((x for x in d["oficial"].get("facturas_normalizadas") or []
                                      if x["numero"] == f["numero"]), None)
                    if oficial_f and (len(oficial_f.get("albaranes", [])), len(oficial_f.get("movimientos", []))) == \
                            (len(f.get("albaranes", [])), len(f.get("movimientos", []))):
                        continue  # sin cambio respecto al oficial
                estado, razon = elegibilidad(f)
                sim = simular_conciliacion(f, operacionales)
                fila = {"origen": origen, "documento": Path(d["ruta"]).name, "numero": f["numero"], "total": f.get("total"),
                        "naturaleza": f.get("naturaleza_principal"), "elegibilidad": [estado, razon], **sim,
                        "produccion": (produccion_numeros.get(f["numero"]) or {}).get("estado_conciliacion_cf")}
                if sim.get("camino") == "PERSISTIR":
                    dif, casados = _d(sim["diferencia"]), sim["albaranes_casados"]
                    tol = tolerancias(casados)
                    fila["estado_por_tolerancia"] = {
                        t: ("CONCILIADA" if abs(dif) <= v else "DIFERENCIA") if estado == "APTA" else f"NO_APTA:{razon}"
                        for t, v in tol.items()}
                    # Variante: la misma regla aplicada tambien al control documental de elegibilidad.
                    tol_doc = tolerancias(len(f.get("albaranes", [])))
                    fila["estado_por_tolerancia_ampliada"] = {}
                    for t in tol:
                        e2, r2 = elegibilidad(f, tol_doc[t])
                        fila["estado_por_tolerancia_ampliada"][t] = (
                            ("CONCILIADA" if abs(dif) <= tol[t] else "DIFERENCIA") if e2 == "APTA" else f"NO_APTA:{r2}")
                    puntos.append((casados, float(dif), f'{f["numero"]} ({origen})'))
                simulaciones.append(fila)
    resumen_t = {}
    for t in ("T1", "T2", "T3"):
        for variante in ("estado_por_tolerancia", "estado_por_tolerancia_ampliada"):
            estados = Counter(s[variante][t].split(":")[0] for s in simulaciones if variante in s)
            cambian = [s["numero"] + f' ({s["origen"]})' for s in simulaciones if variante in s
                       and s[variante][t] != s["estado_por_tolerancia"]["T1"]]
            resumen_t[f"{t}:{variante}"] = {"estados": dict(estados), "cambian_respecto_a_T1": cambian}
    m["conciliacion_simulada"] = {"facturas": simulaciones, "por_tolerancia": resumen_t}
    _svg_dispersion(puntos, salida / "diferencias_vs_albaranes.svg")

    (salida / "metricas.json").write_text(json.dumps(m, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    _muestra(docs, inv, simulaciones, salida / "muestra_contraste.xlsx")
    _tablas(m, salida / "tablas.md")
    for nombre in ("metricas.json", "tablas.md", "diferencias_vs_albaranes.svg", "muestra_contraste.xlsx"):
        print(hashlib.sha256((salida / nombre).read_bytes()).hexdigest(), nombre)


def _muestra(docs, inv, simulaciones, destino: Path) -> None:
    """20 facturas estratificadas: 7 autorizables, 7 no autorizables y 6 simuladas 2AU (solo PIO)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    sims = {(s["documento"], s["numero"], s["origen"]): s for s in simulaciones}
    autorizables, no_autorizables, simuladas = [], [], []
    for d in sorted(docs, key=lambda x: (_proveedor(x), x.get("ruta") or "")):
        if not d.get("ruta") or Path(d["ruta"]).stem.rstrip().casefold().endswith("rita"):
            continue
        o, nombre = d["oficial"], Path(d["ruta"]).name
        normalizadas = o.get("facturas_normalizadas") or []
        for f in normalizadas:
            s = sims.get((nombre, f["numero"], "oficial"))
            res = (f"AUTORIZABLE; elegibilidad {s['elegibilidad'][1]}; conciliación T1 "
                   f"{s.get('resultado_t1')} dif {s.get('diferencia')}" if s else
                   "NO AUTORIZABLE: " + ", ".join(f.get("motivos_no_autorizable") or []))
            fila = [nombre, inv[d["id"]].get("paginas"), f"{o.get('layout')} · {f.get('proveedor_nombre')}", f["numero"],
                    f.get("fecha"), f.get("base_imponible"), f.get("iva"), f.get("total"), res]
            (autorizables if f.get("autorizable") else no_autorizables).append(fila)
        if not normalizadas:
            for f in o.get("facturas_local") or []:
                res = f"NO AUTORIZABLE: {o.get('etapa_detencion')}:{o.get('motivo')}; cuadre " \
                      f"{(f.get('cuadre_fiscal') or {}).get('diferencia')}"
                no_autorizables.append([nombre, inv[d["id"]].get("paginas"), f"{o.get('layout')} · {f.get('proveedor_nombre')}",
                                        f.get("numero"), f.get("fecha"), f.get("base"), f.get("iva"), f.get("total"), res])
        for f in (d.get("simulado_2au") or {}).get("facturas_normalizadas") or []:
            s = sims.get((nombre, f["numero"], "simulado_2au"))
            if s:
                simuladas.append([nombre, inv[d["id"]].get("paginas"), f"{o.get('layout')} · {f.get('proveedor_nombre')}",
                                  f["numero"], f.get("fecha"), f.get("base_imponible"), f.get("iva"), f.get("total"),
                                  f"SIMULACIÓN 2AU: elegibilidad {s['elegibilidad'][1]}; conciliación T1 "
                                  f"{s.get('resultado_t1')} dif {s.get('diferencia')}"])

    def repartir(filas, n):
        if len(filas) <= n:
            return filas
        paso = len(filas) / n
        return [filas[int(i * paso)] for i in range(n)]

    seleccion = ([["AUTORIZABLE", *f] for f in repartir(autorizables, 7)]
                 + [["NO_AUTORIZABLE", *f] for f in repartir(no_autorizables, 7)]
                 + [["SIMULADA_2AU", *f] for f in repartir(simuladas, 6)])
    faltan = 20 - len(seleccion)
    if faltan > 0:
        usados = {tuple(x[1:]) for x in seleccion}
        resto = [["NO_AUTORIZABLE", *f] for f in no_autorizables if tuple(f) not in usados]
        seleccion += repartir(resto, faltan)
    wb = Workbook()
    ws = wb.active
    ws.title = "Muestra 2AY"
    cab = ["Estrato", "Archivo", "Páginas", "Proveedor (layout · nombre)", "Número", "Fecha", "Base", "IVA", "Total",
           "Resultado del motor", "Número REAL", "Fecha REAL", "Base REAL", "IVA REAL", "Total REAL", "¿Coincide?",
           "Observaciones de Pío"]
    ws.append(cab)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for celda in ws[1][10:]:
        celda.fill = PatternFill("solid", fgColor="FFF2CC")
    for fila in seleccion:
        ws.append([*fila, None, None, None, None, None, None, None])
    for col, ancho in zip("ABCDEFGHIJKLMNOPQ", (16, 44, 8, 44, 16, 12, 12, 12, 12, 70, 14, 12, 12, 12, 12, 11, 40)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "B2"
    wb.save(destino)


def _tablas(m: dict, destino: Path) -> None:
    partes = ["## Embudo global", _tabla(["Etapa", "Valor"], [[k, v] for k, v in m["embudo"].items()]),
              "## Etapa de detención (documentos)", _tabla(["Etapa:motivo", "Docs"], m["etapas_detencion"].most_common()
                                                           if isinstance(m["etapas_detencion"], Counter)
                                                           else sorted(m["etapas_detencion"].items(), key=lambda x: -x[1])),
              "## Por proveedor", _tabla(["Layout", "Docs", "Completos", "Facturas", "Cuadre fiscal", "Autorizables", "%",
                                          "3 motivos principales"],
                                         [[k, v["documentos"], v["completos"], v["facturas"], v["facturas_con_cuadre_fiscal"],
                                           v["autorizables"], v["pct_autorizables"],
                                           "; ".join(f"{a} ({b})" for a, b in v["motivos_principales"])]
                                          for k, v in sorted(m["por_proveedor"].items(), key=lambda x: -x[1]["documentos"])])]
    destino.write_text("\n\n".join(partes) + "\n", encoding="utf-8")


if __name__ == "__main__":
    destino = Path(sys.argv[4]).resolve()
    if destino == ROOT or ROOT in destino.parents:
        raise SystemExit("SALIDA_DENTRO_DEL_REPOSITORIO")
    destino.mkdir(parents=True, exist_ok=True)
    main(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), destino)
